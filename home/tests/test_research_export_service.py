"""
Bulk export of consented research data (shared surveys, participants who gave research consent).
"""

import io
import tempfile
from io import StringIO
from pathlib import Path
from unittest import mock

import openpyxl
from django.core.management import CommandError, call_command
from django.test import TestCase

from home.services.research_export import (
    EXAMPLES_PER_SURVEY,
    METADATA_COLUMNS,
    QUESTIONS_COLUMNS,
    QUESTIONS_SHEET_NAME,
    SHEET_NAME,
    ResearchExport,
)
from SORT.test.model_factory import SurveyFactory
from survey.models import RESEARCH_CONSENT_FIELD_NAME, Profession, Survey, SurveyResponse

# A custom field added in the config editor has no name, so its label identifies its column
EXTRA_QUESTION = "Extra question"
# The SORT likert sub-questions (A1 to E9) and the comments field of each of the 5 SORT sections
SORT_CODE_COUNT = 44
SORT_COMMENTS_COUNT = 5


def field_index(survey: Survey, name: str) -> int:
    fields = survey.sections[0]["fields"]
    return next(i for i, f in enumerate(fields) if f["name"] == name)


def make_survey(
    is_shared: bool = True, extra_question: bool = False, profession: Profession = Profession.NURSES
) -> Survey:
    survey = SurveyFactory(is_shared=is_shared, survey_body_path=profession)
    survey.initialise()
    if extra_question:
        survey.survey_config["sections"].append(
            {"title": "Extra", "fields": [{"type": "text", "label": EXTRA_QUESTION}]}
        )
    survey.save()
    return survey


def make_response(survey: Survey, research_consent: bool, extra_answer: str = None) -> SurveyResponse:
    answers = survey._generate_mock_response()
    index = field_index(survey, RESEARCH_CONSENT_FIELD_NAME)
    option = survey.sections[0]["fields"][index]["options"][0]
    answers[0][index] = [option] if research_consent else []
    if extra_answer is not None:
        answers[-1] = [extra_answer]
    return SurveyResponse.objects.create(survey=survey, answers=answers)


def make_duplicate_label_survey() -> Survey:
    """A survey whose customised section has two fields with the same label (as in production survey 180)."""
    survey = make_survey()
    survey.survey_config["sections"].append(
        {
            "title": "Extra",
            "fields": [{"type": "text", "label": EXTRA_QUESTION}, {"type": "text", "label": EXTRA_QUESTION}],
        }
    )
    survey.save()
    return survey


def read_sheet(export: ResearchExport, sheet_name: str = SHEET_NAME) -> list[tuple]:
    buffer = io.BytesIO()
    export.write_xlsx(buffer)
    workbook = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    return list(workbook[sheet_name].iter_rows(values_only=True))


class ResearchExportTestCase(TestCase):
    def test_includes_only_consented_responses_on_shared_surveys(self):
        shared = make_survey(is_shared=True)
        unshared = make_survey(is_shared=False)
        consented = make_response(shared, research_consent=True)
        make_response(shared, research_consent=False)
        make_response(unshared, research_consent=True)

        rows = read_sheet(ResearchExport())
        response_id_col = METADATA_COLUMNS.index("Response ID")
        self.assertEqual([row[response_id_col] for row in rows[1:]], [consented.pk])

    def test_metadata_columns(self):
        survey = make_survey()
        response = make_response(survey, research_consent=True)

        header, row = read_sheet(ResearchExport())
        self.assertEqual(header[: len(METADATA_COLUMNS)], METADATA_COLUMNS)
        self.assertEqual(
            row[:5],
            (survey.organisation.name, survey.project.name, survey.pk, survey.name, response.pk),
        )
        self.assertIsNotNone(row[5])

    def test_shared_survey_without_project(self):
        survey = make_survey()
        survey.project = None
        survey.save()
        response = make_response(survey, research_consent=True)

        _, row = read_sheet(ResearchExport())
        self.assertEqual(row[:5], (None, None, survey.pk, survey.name, response.pk))

    def test_skips_shared_survey_without_config(self):
        uninitialised = make_survey()
        make_response(uninitialised, research_consent=True)
        Survey.objects.filter(pk=uninitialised.pk).update(survey_config=None)
        valid = make_survey()
        make_response(valid, research_consent=True)

        rows = read_sheet(ResearchExport())
        self.assertEqual(len(rows), 2)

    def test_columns_are_union_of_survey_questions(self):
        plain = make_survey()
        extended = make_survey(extra_question=True)
        make_response(plain, research_consent=True)
        make_response(extended, research_consent=True, extra_answer="Hello")

        header, plain_row, extended_row = read_sheet(ResearchExport())
        self.assertEqual(header.count(EXTRA_QUESTION), 1)
        self.assertEqual(header[len(METADATA_COLUMNS) :], (*plain.field_ids, EXTRA_QUESTION))
        extra_col = header.index(EXTRA_QUESTION)
        self.assertIsNone(plain_row[extra_col])
        self.assertEqual(extended_row[extra_col], "Hello")

    def test_profession_variants_share_question_columns(self):
        """
        The same SORT question is worded differently in each profession variant, but shares a column.
        See https://github.com/RSE-Sheffield/SORT/issues/799
        """
        nurses = make_survey(profession=Profession.NURSES)
        ahp = make_survey(profession=Profession.AHP)
        self.assertNotEqual(nurses.fields, ahp.fields)
        make_response(nurses, research_consent=True)
        make_response(ahp, research_consent=True)

        header, *rows = read_sheet(ResearchExport())
        codes = [column for column in header if column[0] in "ABCDE" and column[1:].isdigit()]
        comments = [column for column in header if column.endswith("_comments")]
        self.assertEqual(len(codes), SORT_CODE_COUNT)
        self.assertEqual(len(comments), SORT_COMMENTS_COUNT)
        self.assertEqual(len(header), len(METADATA_COLUMNS) + len(nurses.field_ids))
        for row in rows:
            self.assertTrue(all(row[header.index(code)] for code in codes))

    def test_duplicate_labels_are_separate_columns(self):
        survey = make_duplicate_label_survey()
        response = make_response(survey, research_consent=True)
        response.answers[-1] = ["First", "Second"]
        response.save()

        header, row = read_sheet(ResearchExport())
        self.assertEqual(header[-2:], (EXTRA_QUESTION, f"{EXTRA_QUESTION} (2)"))
        self.assertEqual(row[-2:], ("First", "Second"))

    def test_questions_sheet_maps_columns_to_each_wording(self):
        nurses = make_survey(profession=Profession.NURSES)
        ahp = make_survey(profession=Profession.AHP)
        make_response(nurses, research_consent=True)
        make_response(ahp, research_consent=True)

        header, *rows = read_sheet(ResearchExport(), QUESTIONS_SHEET_NAME)
        self.assertEqual(header, QUESTIONS_COLUMNS)
        wordings = {(question, survey_ids) for column, question, survey_ids in rows if column == "A2"}
        a2 = nurses.field_ids.index("A2")
        self.assertEqual(wordings, {(nurses.fields[a2], str(nurses.pk)), (ahp.fields[a2], str(ahp.pk))})
        # A question worded the same in both surveys has one row listing both
        self.assertIn(("consent-research", ahp.fields[1], f"{nurses.pk}, {ahp.pk}"), rows)

    def test_surveys_without_consented_responses_add_no_columns(self):
        plain = make_survey()
        extended = make_survey(extra_question=True)
        make_response(plain, research_consent=True)
        make_response(extended, research_consent=False, extra_answer="Hello")

        header, _ = read_sheet(ResearchExport())
        self.assertNotIn(EXTRA_QUESTION, header)

    def test_answers_are_not_written_as_formulas(self):
        survey = make_survey(extra_question=True)
        make_response(survey, research_consent=True, extra_answer="=1+1")

        header, row = read_sheet(ResearchExport())
        self.assertEqual(row[header.index(EXTRA_QUESTION)], "=1+1")
        buffer = io.BytesIO()
        ResearchExport().write_xlsx(buffer)
        sheet = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))[SHEET_NAME]
        self.assertEqual(sheet.cell(row=2, column=header.index(EXTRA_QUESTION) + 1).data_type, "s")

    def test_empty_export_has_header_only(self):
        rows = read_sheet(ResearchExport())
        self.assertEqual(rows, [METADATA_COLUMNS])


class ResearchExportDiagnoseTestCase(TestCase):
    def test_valid_responses_have_no_malformed_blanks(self):
        survey = make_survey()
        make_response(survey, research_consent=True)

        diagnosis = ResearchExport().diagnose()
        self.assertEqual(diagnosis.row_count, 1)
        self.assertEqual(diagnosis.column_count, len(survey.field_ids))
        self.assertEqual(diagnosis.blank_other_survey, 0)
        self.assertEqual(diagnosis.blank_not_stored, 0)
        self.assertEqual(diagnosis.short_responses, {})
        self.assertEqual(diagnosis.invalid_responses, {})

    def test_counts_blanks_from_other_surveys_questions(self):
        plain = make_survey()
        extended = make_survey(extra_question=True)
        make_response(plain, research_consent=True)
        make_response(extended, research_consent=True, extra_answer="Hello")

        diagnosis = ResearchExport().diagnose()
        # The plain survey's row has no answer for the extended survey's extra question
        self.assertEqual(diagnosis.blank_other_survey, 1)
        self.assertEqual(diagnosis.blank_not_stored, 0)
        self.assertEqual(diagnosis.rows_per_survey, {plain.pk: 1, extended.pk: 1})

    def test_counts_answers_missing_from_stored_json(self):
        survey = make_survey()
        response = make_response(survey, research_consent=True)
        response.answers[-1] = []  # e.g. demography saved before response schema validation
        response.save()
        missing = len(survey.sections[-1]["fields"])

        diagnosis = ResearchExport().diagnose()
        self.assertEqual(diagnosis.blank_not_stored, missing)
        self.assertEqual(diagnosis.short_responses, {survey.pk: 1})
        self.assertEqual(diagnosis.invalid_responses, {survey.pk: 1})
        self.assertEqual([pk for pk, _ in diagnosis.invalid_examples[survey.pk]], [response.pk])

    def test_keeps_only_first_invalid_examples(self):
        survey = make_survey()
        responses = [make_response(survey, research_consent=True) for _ in range(3)]
        for response in responses:
            response.answers[-1] = []
            response.save()

        diagnosis = ResearchExport().diagnose(max_examples=2)
        self.assertEqual(diagnosis.invalid_responses, {survey.pk: 3})
        self.assertEqual([pk for pk, _ in diagnosis.invalid_examples[survey.pk]], [r.pk for r in responses[:2]])

    def test_reports_duplicate_labels_without_losing_answers(self):
        survey = make_duplicate_label_survey()
        response = make_response(survey, research_consent=True)
        response.answers[-1] = ["First", ""]
        response.save()

        diagnosis = ResearchExport().diagnose()
        self.assertEqual(diagnosis.duplicate_labels, {survey.pk: {EXTRA_QUESTION: 2}})
        self.assertEqual(diagnosis.column_count, len(survey.field_ids))
        # Only the second field is blank: the first answer is not overwritten
        self.assertEqual(diagnosis.blank_other_survey, 0)
        self.assertEqual(diagnosis.blank_not_stored, 0)
        self.assertEqual(diagnosis.blank_count, diagnosis.blank_empty)

    def test_builds_one_validator_per_survey(self):
        surveys = [make_survey(), make_survey()]
        for survey in surveys:
            for _ in range(3):
                make_response(survey, research_consent=True)

        # Both surveys share the default configuration, so one real validator serves both
        validator = surveys[0].response_validator
        with mock.patch.object(
            Survey, "response_validator", new_callable=mock.PropertyMock, return_value=validator
        ) as response_validator:
            diagnosis = ResearchExport().diagnose()
        self.assertEqual(diagnosis.row_count, 6)
        self.assertEqual(diagnosis.invalid_responses, {})
        self.assertEqual(response_validator.call_count, len(surveys))

    def test_answers_stored_aligns_with_answers_values(self):
        survey = make_survey()
        response = make_response(survey, research_consent=True)
        response.answers[-1] = []
        stored = list(response.answers_stored)

        self.assertEqual(len(stored), len(list(response.answers_values)))
        self.assertEqual(len(stored), len(survey.fields))
        self.assertEqual(stored.count(False), len(survey.sections[-1]["fields"]))


class ResearchExportCommandTestCase(TestCase):
    def test_writes_workbook(self):
        survey = make_survey()
        make_response(survey, research_consent=True)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "research.xlsx"
            out = StringIO()
            call_command("research_export", "-o", str(path), stdout=out)
            self.assertIn("Wrote 1 responses", out.getvalue())
            rows = list(openpyxl.load_workbook(path)[SHEET_NAME].iter_rows(values_only=True))
            self.assertEqual(len(rows), 2)

    def test_diagnose(self):
        survey = make_survey()
        response = make_response(survey, research_consent=True)
        response.answers[-1] = []
        response.save()
        out = StringIO()
        call_command("research_export", "--diagnose", stdout=out)
        output = out.getvalue()
        self.assertIn("Rows: 1 from 1 surveys", output)
        self.assertIn(f"Survey {survey.pk}: 1 of 1 responses", output)
        self.assertIn(f"Response {response.pk}:", output)

    def test_requires_output_without_diagnose(self):
        with self.assertRaises(CommandError):
            call_command("research_export", stdout=StringIO())

    def test_rejects_output_with_diagnose(self):
        with self.assertRaises(CommandError):
            call_command("research_export", "--diagnose", "-o", "research.xlsx", stdout=StringIO())

    def test_diagnose_says_how_many_examples_were_omitted(self):
        survey = make_survey()
        for _ in range(EXAMPLES_PER_SURVEY + 2):
            response = make_response(survey, research_consent=True)
            response.answers[-1] = []
            response.save()

        out = StringIO()
        call_command("research_export", "--diagnose", stdout=out)
        self.assertIn("... and 2 more (use -v 2 to list all)", out.getvalue())

        out = StringIO()
        call_command("research_export", "--diagnose", verbosity=2, stdout=out)
        self.assertNotIn("more (use -v 2", out.getvalue())
